# -*- coding: utf-8 -*-
"""
Serial Number Filler (Streamlit)
1) upload Blank xlsx  2) upload Service / PM report scans -> AI reads Date, customer, Serial No., form type
3) preview + edit      4) export xlsx (writes Serial number, Date, Service type only)

Run app :  streamlit run app.py
Self test: python app.py --selftest
"""
from __future__ import annotations

import base64
import copy
import difflib
import hashlib
import io
import json
import os
import posixpath
import re
import sys
import time
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

import pandas as pd
import requests
import streamlit as st
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import column_index_from_string, get_column_letter
from PIL import Image, ImageOps, ImageSequence

try:
    import fitz  # PyMuPDF
except Exception:  # pragma: no cover
    fitz = None

# ----------------------------------------------------------------------------
# Constants
# ----------------------------------------------------------------------------
APP_DIR = Path(__file__).resolve().parent
DATA_DIR = APP_DIR / "data"
STORE_PATH = DATA_DIR / "serials_v2.json"
TEMPLATE_PATH = DATA_DIR / "template_v1.xlsx"

SHEET_SN = "SN tracking form"
SHEET_DATA2 = "Data2 (2)"

PROVIDER_GEMINI = "Google Gemini (ฟรี)"
PROVIDER_CLAUDE = "Anthropic (Claude)"
PROVIDER_OPENAI = "OpenAI"
PROVIDER_DEMO = "Demo (ไม่ใช้ AI - ทดสอบระบบ)"
DEFAULT_MODELS = {PROVIDER_GEMINI: "gemini-3.5-flash-lite", PROVIDER_CLAUDE: "claude-sonnet-4-5",
                  PROVIDER_OPENAI: "gpt-4o", PROVIDER_DEMO: "-"}

PDF_ZOOM = 2.0
MAX_SIDE = 1800
JPEG_QUALITY = 85
GEMINI_PAUSE = 4.0          # seconds between AI calls (free tier rate limit)
COLS = ["serial", "date", "customer", "form", "source", "page"]
FORM_OPTIONS = ["Service", "PM"]
CROP_OPTIONS = {"ครึ่งบน 50% (แนะนำ)": 0.5, "บน 40% (เร็วกว่า)": 0.4,
                "บน 65% (เผื่อฟอร์มเลื่อนลง)": 0.65, "ทั้งหน้า (ช้าสุด)": 1.0}

SYSTEM_PROMPT = (
    "You read the header fields of Nipro (Thailand) dialysis-machine service forms "
    "(Thai + English, handwritten). Return ONLY valid JSON. Never invent characters."
)
USER_PROMPT = """This image is the TOP PART (or the whole page) of a Nipro form: a 'Service Report' or a 'Preventive Maintenance Report'.
Read ONLY these header fields and ignore everything else (parts, comments, tables, signatures, Book No., No., M/C no.):
- form_type: "service" if the title says Service Report; "pm" if it says Preventive Maintenance Report; otherwise "other".
- date: the 'Date' field as YYYY-MM-DD. The day comes first (dd/mm/yy): 6/8/26 -> 2026-08-06. Buddhist-era years: 2569 or 69 -> 2026. Use "" if unreadable.
- customer_name: the hospital / customer name written on the line beside the title (Thai or English, copy as written, "" if unreadable).
- candidates: Service Report has two 'Serial no.' boxes: return the text of BOTH boxes ("" for an empty box). The real serial is the long one.
  Preventive Maintenance Report: the box labelled 'Serial No.' (NOT 'M/C no.').
- serial: the most likely serial number.
Serial format: 2 digits + 2 letters + digits, e.g. 21CX22507, 17CX6407, 22CX26509, 14AS4052; sometimes digits with a dash like 2403-055.
A handwritten 'C' may look like '(' or 'c' and 'X' may look like 'K': if the text clearly matches the pattern, write the standard letters.
If a character is truly unreadable write '?'. Do not invent digits.
If the image holds more than one form return one item per form. If it is not a form return {"reports": []}.
Return JSON exactly like:
{"reports":[{"form_type":"service","date":"2026-08-06","customer_name":"","candidates":["",""],"serial":""}]}"""

DEMO_PAYLOAD = {"reports": [{"form_type": "service", "date": "2026-08-06",
                             "customer_name": "Rajthanee Hospital (Ayutthaya)",
                             "candidates": ["21CX22507", "14"], "serial": "21CX22507"}]}


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def is_blank(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, float) and v != v:
        return True
    return isinstance(v, str) and not v.strip()


def _s(v: Any) -> str:
    return "" if is_blank(v) else str(v).strip()


def norm_key(v: Any) -> str:
    return "" if is_blank(v) else re.sub(r"\s+", " ", str(v)).strip().casefold()


def clean_serial(v: Any) -> str:
    if is_blank(v):
        return ""
    return re.sub(r"[^A-Z0-9?\-]", "", str(v).upper())


def alnum_len(s: str) -> int:
    return len(re.sub(r"[^A-Z0-9]", "", s))


def pick_serial(cands: list) -> str:
    cl = [clean_serial(c) for c in cands if not is_blank(c)]
    cl = [c for c in cl if c]
    if not cl:
        return ""
    return max(cl, key=lambda s: (alnum_len(s), -s.count("?")))


def parse_date(v: Any) -> Optional[date]:
    if is_blank(v):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    else:
        m = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})", s)
        if not m:
            return None
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if y < 100:
        y = y + 2500 - 543 if y >= 50 else y + 2000   # 69 -> BE 2569 -> 2026 ; 26 -> 2026
    elif y > 2400:
        y -= 543
    try:
        return date(y, mo, d)
    except ValueError:
        return None


def parse_json_loose(text: str) -> dict:
    t = (text or "").strip()
    t = re.sub(r"^```(?:json)?\s*", "", t, flags=re.I)
    t = re.sub(r"\s*```$", "", t)
    try:
        return json.loads(t)
    except Exception:
        pass
    s, e = t.find("{"), t.rfind("}")
    if s != -1 and e > s:
        return json.loads(t[s:e + 1])
    raise ValueError("AI response is not valid JSON")


# ----------------------------------------------------------------------------
# Storage (so a browser refresh does not lose the work)
# ----------------------------------------------------------------------------
def load_store() -> dict:
    try:
        d = json.loads(STORE_PATH.read_text(encoding="utf-8"))
        return {"rows": list(d.get("rows", [])), "done": list(d.get("done", []))}
    except Exception:
        return {"rows": [], "done": []}


def save_store(rows: list[dict], done) -> None:
    try:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        STORE_PATH.write_text(json.dumps({"rows": rows, "done": sorted(done)}, ensure_ascii=False),
                              encoding="utf-8")
    except Exception:
        pass


def df_to_rows(df: pd.DataFrame) -> list[dict]:
    out = []
    for r in df.to_dict("records"):
        serial = clean_serial(r.get("serial"))
        src, raw_d, cust = _s(r.get("source")), _s(r.get("date")), _s(r.get("customer"))
        if not (serial or src or raw_d or cust):
            continue
        d = parse_date(raw_d)
        form = _s(r.get("form"))
        out.append({"serial": serial, "date": d.isoformat() if d else raw_d, "customer": cust,
                    "form": form if form in FORM_OPTIONS else "", "source": src, "page": _s(r.get("page"))})
    return out


# ----------------------------------------------------------------------------
# Page rendering (PDF / images -> JPEG, optionally only the top part)
# ----------------------------------------------------------------------------
def _to_jpeg(img: Image.Image, crop: float = 1.0) -> bytes:
    img = ImageOps.exif_transpose(img).convert("RGB")
    if 0 < crop < 1:
        w, h = img.size
        img = img.crop((0, 0, w, max(int(h * crop), 1)))
    if max(img.size) > MAX_SIDE:
        img.thumbnail((MAX_SIDE, MAX_SIDE), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=JPEG_QUALITY)
    return buf.getvalue()


def _pdf_page_image(page) -> Image.Image:
    pix = page.get_pixmap(matrix=fitz.Matrix(PDF_ZOOM, PDF_ZOOM), colorspace=fitz.csRGB, alpha=False)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def count_pages(name: str, data: bytes) -> int:
    if name.lower().endswith(".pdf"):
        if fitz is None:
            raise RuntimeError("ไม่พบไลบรารี PyMuPDF")
        doc = fitz.open(stream=data, filetype="pdf")
        try:
            if doc.needs_pass:
                raise RuntimeError("PDF ถูกตั้งรหัสผ่าน")
            return doc.page_count
        finally:
            doc.close()
    with Image.open(io.BytesIO(data)) as im:
        return max(getattr(im, "n_frames", 1), 1)


def iter_pages(name: str, data: bytes, crop: float = 1.0):
    if name.lower().endswith(".pdf"):
        doc = fitz.open(stream=data, filetype="pdf")
        try:
            for i in range(doc.page_count):
                yield i + 1, _to_jpeg(_pdf_page_image(doc[i]), crop)
        finally:
            doc.close()
    else:
        with Image.open(io.BytesIO(data)) as im:
            for i, frame in enumerate(ImageSequence.Iterator(im)):
                yield i + 1, _to_jpeg(frame.copy(), crop)


def render_page(name: str, data: bytes, page_no: int, crop: float = 1.0) -> bytes:
    if name.lower().endswith(".pdf"):
        doc = fitz.open(stream=data, filetype="pdf")
        try:
            return _to_jpeg(_pdf_page_image(doc[page_no - 1]), crop)
        finally:
            doc.close()
    with Image.open(io.BytesIO(data)) as im:
        im.seek(page_no - 1)
        return _to_jpeg(im.copy(), crop)


# ----------------------------------------------------------------------------
# AI client
# ----------------------------------------------------------------------------
class AiFatal(Exception):
    """Retrying will not help (bad key / model / quota)."""


class AiQuota(AiFatal):
    pass


def _post(url: str, headers: dict, payload: dict) -> dict:
    r = requests.post(url, headers=headers, json=payload, timeout=120)
    if r.status_code == 200:
        return r.json()
    msg = f"HTTP {r.status_code}: {r.text[:300]}"
    if r.status_code == 429:
        raise AiQuota(msg)
    if r.status_code in (400, 401, 403, 404):
        raise AiFatal(msg)
    raise RuntimeError(msg)


class VisionClient:
    def __init__(self, provider: str, api_key: str, model: str, pause: float = 0.0) -> None:
        self.provider, self.api_key, self.model = provider, api_key.strip(), model.strip()
        self.pause = float(pause or 0.0)

    def extract(self, jpeg: bytes) -> dict:
        if self.provider == PROVIDER_DEMO:
            return json.loads(json.dumps(DEMO_PAYLOAD))
        if not self.api_key:
            raise AiFatal("ยังไม่ได้ใส่ API key")
        waits = [4, 12]
        quota_retry = 1
        while True:
            try:
                return parse_json_loose(self._call(jpeg))
            except AiQuota:
                if quota_retry:
                    quota_retry -= 1
                    time.sleep(30)          # a per-minute limit may clear; a daily limit will fail again
                    continue
                raise
            except AiFatal:
                raise
            except Exception as e:
                if waits:
                    time.sleep(waits.pop(0))
                    continue
                raise RuntimeError(f"AI extraction failed: {e}")

    def _call(self, jpeg: bytes) -> str:
        b64 = base64.b64encode(jpeg).decode("ascii")
        if self.provider == PROVIDER_GEMINI:
            gen: dict[str, Any] = {"responseMimeType": "application/json", "temperature": 0}
            if "2.5" in self.model and "pro" not in self.model:
                gen["thinkingConfig"] = {"thinkingBudget": 0}
            data = _post(
                f"https://generativelanguage.googleapis.com/v1beta/models/{self.model}:generateContent",
                {"x-goog-api-key": self.api_key, "Content-Type": "application/json"},
                {"systemInstruction": {"parts": [{"text": SYSTEM_PROMPT}]},
                 "contents": [{"role": "user", "parts": [{"text": USER_PROMPT},
                                                          {"inlineData": {"mimeType": "image/jpeg", "data": b64}}]}],
                 "generationConfig": gen})
            cands = data.get("candidates") or []
            if not cands:
                raise RuntimeError(f"Gemini ไม่ตอบ: {data.get('promptFeedback')}")
            return "".join(p.get("text", "") for p in (cands[0].get("content") or {}).get("parts") or [])
        if self.provider == PROVIDER_CLAUDE:
            data = _post(
                "https://api.anthropic.com/v1/messages",
                {"x-api-key": self.api_key, "anthropic-version": "2023-06-01", "content-type": "application/json"},
                {"model": self.model, "max_tokens": 1024, "system": SYSTEM_PROMPT,
                 "messages": [{"role": "user", "content": [
                     {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": b64}},
                     {"type": "text", "text": USER_PROMPT}]}]})
            return "".join(b.get("text", "") for b in data.get("content", []))
        if self.provider == PROVIDER_OPENAI:
            data = _post(
                "https://api.openai.com/v1/chat/completions",
                {"Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json"},
                {"model": self.model, "response_format": {"type": "json_object"},
                 "messages": [{"role": "system", "content": SYSTEM_PROMPT},
                              {"role": "user", "content": [
                                  {"type": "text", "text": USER_PROMPT},
                                  {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}}]}]})
            return data["choices"][0]["message"]["content"] or ""
        raise AiFatal(f"Unknown provider: {self.provider}")


def reports_to_rows(raw: Any, source: str, page: int) -> list[dict]:
    items = raw.get("reports") if isinstance(raw, dict) else None
    if not isinstance(items, list):
        items = []
    rows = []
    for it in items:
        if not isinstance(it, dict):
            continue
        cands = it.get("candidates") or []
        if isinstance(cands, str):
            cands = [cands]
        serial = pick_serial(list(cands) + [it.get("serial")])
        form = {"service": "Service", "pm": "PM"}.get(str(it.get("form_type", "")).strip().lower(), "")
        raw_d = _s(it.get("date"))
        d = parse_date(raw_d)
        rows.append({"serial": serial, "date": d.isoformat() if d else raw_d,
                     "customer": _s(it.get("customer_name")), "form": form,
                     "source": source, "page": str(page)})
    if not rows:   # blank / unreadable page -> visible placeholder row so nothing is silently lost
        rows.append({"serial": "", "date": "", "customer": "", "form": "", "source": source, "page": str(page)})
    return rows


def read_page(client: VisionClient, jpeg: bytes, name: str, data: bytes, page_no: int, crop: float) -> list[dict]:
    new = reports_to_rows(client.extract(jpeg), name, page_no)
    if client.pause:
        time.sleep(client.pause)
    if crop < 1 and not any(r["serial"] for r in new):      # nothing found in the top part -> try the whole page
        try:
            full = render_page(name, data, page_no, 1.0)
            new2 = reports_to_rows(client.extract(full), name, page_no)
            if client.pause:
                time.sleep(client.pause)
            if any(r["serial"] for r in new2):
                new = new2
        except AiFatal:
            raise
        except Exception:
            pass
    return new


def process_files(files: list[tuple[str, bytes]], client: VisionClient, rows: list[dict], done: set,
                  crop: float = 0.5, on_progress=None) -> dict:
    res = {"pages": 0, "added": 0, "skipped": 0, "errors": [], "stopped": ""}
    plan = []
    for name, data in files:
        try:
            plan.append((name, data, count_pages(name, data)))
        except Exception as e:
            res["errors"].append(f"{name}: เปิดไฟล์ไม่ได้ ({e})")
    total = max(sum(p[2] for p in plan), 1)
    seen = 0
    fails = 0
    for name, data, _n in plan:
        fhash = hashlib.sha1(data).hexdigest()[:16]
        try:
            for page_no, jpeg in iter_pages(name, data, crop):
                seen += 1
                key = f"{fhash}:{page_no}"
                if on_progress:
                    on_progress(seen / total, f"{name} หน้า {page_no}")
                if key in done:
                    res["skipped"] += 1
                    continue
                try:
                    new = read_page(client, jpeg, name, data, page_no, crop)
                except AiFatal as e:
                    res["stopped"] = str(e)[:300]
                    return res
                except Exception as e:
                    res["errors"].append(f"{name} หน้า {page_no}: {str(e)[:200]}")
                    fails += 1
                    if fails >= 3:
                        res["stopped"] = "AI ตอบผิดพลาดติดกัน 3 หน้า (อาจเป็นเซิร์ฟเวอร์ Google คนใช้เยอะ ลองใหม่ภายหลัง)"
                        return res
                    continue
                fails = 0
                rows.extend(new)
                done.add(key)
                save_store(rows, done)
                res["pages"] += 1
                res["added"] += len(new)
        except Exception as e:
            res["errors"].append(f"{name}: {e}")
    return res


# ----------------------------------------------------------------------------
# Reference check against the Data2 sheet of the uploaded workbook
# ----------------------------------------------------------------------------
def ref_serials_impl(data: bytes) -> dict[str, dict]:
    out: dict[str, dict] = {}
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except Exception:
        return out
    try:
        ws = next((w for w in wb.worksheets if norm_key(w.title) == norm_key(SHEET_DATA2)), None)
        if ws is None:
            return out
        idx = en_i = th_i = None
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if idx is None:
                keys = [norm_key(c) for c in row]
                if "serial number" in keys:
                    idx = keys.index("serial number")
                    en_i = next((j for j, k in enumerate(keys) if k.startswith("eng - customer name")), None)
                    th_i = next((j for j, k in enumerate(keys) if k.startswith("th - customer name")), None)
                elif i > 15:
                    break
                continue
            if idx < len(row):
                s = clean_serial(row[idx])
                if s and s not in out:
                    en = _s(row[en_i]) if en_i is not None and en_i < len(row) else ""
                    th = _s(row[th_i]) if th_i is not None and th_i < len(row) else ""
                    out[s] = {"raw": str(row[idx]).strip(), "name": en or th}
    finally:
        wb.close()
    return out


def check_serial(s: str, ref: dict, ref_keys: list[str]) -> tuple[str, str]:
    if not ref:
        return "", ""
    if not s:
        return "ว่าง", ""
    if s in ref:
        return "พบ", ""
    cm = difflib.get_close_matches(s, ref_keys, n=1, cutoff=0.8)
    return "ไม่พบ", (ref[cm[0]]["raw"] if cm else "")


# ----------------------------------------------------------------------------
# Locate the Serial number / Date / Service type columns in the template
# ----------------------------------------------------------------------------
def _has_style(c) -> bool:
    try:
        if c.fill is not None and c.fill.fill_type:
            return True
        b = c.border
        return any(getattr(getattr(b, side, None), "style", None) for side in ("left", "right", "top", "bottom"))
    except Exception:
        return False


def _is_const(v: Any) -> bool:
    return not is_blank(v) and not (isinstance(v, str) and v.startswith("="))


def _find_in_header(ws, hdr: int, pred) -> Optional[int]:
    for r in (hdr, hdr + 1, hdr - 1):
        if r < 1:
            continue
        for c in range(1, 41):
            if pred(norm_key(ws.cell(r, c).value)):
                return c
    return None


def analyse_sheet(wb) -> dict:
    ws = next((w for w in wb.worksheets if norm_key(w.title) == norm_key(SHEET_SN)), None)
    if ws is None:
        return {"ok": False, "error": f"ไม่พบชีต '{SHEET_SN}' ในไฟล์"}
    hdr = col = None
    for r in range(1, 41):
        for c in range(1, 41):
            if norm_key(ws.cell(r, c).value) == "serial number":
                hdr, col = r, c
                break
        if hdr:
            break
    if hdr is None:
        return {"ok": False, "error": "ไม่พบหัวคอลัมน์ 'Serial number' ในชีต SN tracking form"}
    date_col = _find_in_header(ws, hdr, lambda t: t == "date")
    type_col = _find_in_header(ws, hdr, lambda t: t.startswith("service type"))
    first = hdr + 1
    for c in range(1, 12):
        t = norm_key(ws.cell(first, c).value)
        if "(only name)" in t or "(drop down" in t:
            first += 1
            break
    while type(ws.cell(first, col)).__name__ == "MergedCell" and first < hdr + 6:
        first += 1
    limit = min(ws.max_row, first + 3000)
    stop = None
    for r in range(first, limit):
        for c in range(1, 13):
            if norm_key(ws.cell(r, c).value) == "warranty" and norm_key(ws.cell(r + 1, c).value) == "contract":
                stop = r
                break
        if stop:
            break
    end = (stop - 1) if stop else limit
    last_used = first - 1
    existing = 0
    src = None
    for r in range(first, end + 1):
        cell = ws.cell(r, col)
        if _is_const(cell.value):
            last_used, existing = r, existing + 1
        if _has_style(cell):
            src = r
    start = last_used + 1
    free = None
    if stop:
        src = max(src or (stop - 1), last_used)
        free = max(src - start + 1, 0)
    return {"ok": True, "title": ws.title, "hdr": hdr, "col": col, "col_letter": get_column_letter(col),
            "date_col": date_col, "date_letter": get_column_letter(date_col) if date_col else None,
            "type_col": type_col, "type_letter": get_column_letter(type_col) if type_col else None,
            "first": first, "start": start, "stop": stop, "src": src, "free": free, "existing": existing}


def plan_template_impl(data: bytes) -> dict:
    try:
        wb = load_workbook(io.BytesIO(data), keep_links=False)
    except Exception as e:
        return {"ok": False, "error": f"เปิดไฟล์ไม่ได้: {e}"}
    try:
        return analyse_sheet(wb)
    finally:
        wb.close()


# ----------------------------------------------------------------------------
# XML-level patching (keeps formulas, styles, drop downs, links, shapes of the original file)
# ----------------------------------------------------------------------------
_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_XML_NS = "http://www.w3.org/XML/1998/namespace"
_REF_RE = re.compile(
    r"(?<![\w.])(?P<pre>(?:'[^']+'|[A-Za-z_][\w.]*)!)?"
    r"(?P<c1>\$?[A-Z]{1,3})(?P<r1>\$?)(?P<n1>\d+)"
    r"(?::(?P<c2>\$?[A-Z]{1,3})(?P<r2>\$?)(?P<n2>\d+))?"
    r"(?![\w(!])"
)


def _q(tag: str) -> str:
    return f"{{{_NS}}}{tag}"


class DateCell:
    """A real Excel date (serial number) to write into a cell."""
    __slots__ = ("serial",)

    def __init__(self, serial: float) -> None:
        self.serial = float(serial)


class StyleBook:
    """Gives a date number format to a cell style without touching its fill / border / font."""

    def __init__(self, xml: bytes) -> None:
        self.root = etree.fromstring(xml)
        self.xfs = self.root.find(_q("cellXfs"))
        self.numfmts = self.root.find(_q("numFmts"))
        self.custom: dict[int, str] = {}
        if self.numfmts is not None:
            for nf in self.numfmts.findall(_q("numFmt")):
                try:
                    self.custom[int(nf.get("numFmtId"))] = nf.get("formatCode") or ""
                except (TypeError, ValueError):
                    pass
        self.fmt_id: Optional[int] = None
        self.cache: dict[int, int] = {}

    def _is_date(self, fid: int) -> bool:
        if fid in (14, 15, 16, 17, 18, 19, 20, 21, 22, 45, 46, 47):
            return True
        code = self.custom.get(fid)
        return bool(code) and is_date_format(code)

    def _add_fmt(self) -> int:
        if self.numfmts is None:
            self.numfmts = etree.Element(_q("numFmts"))
            self.root.insert(0, self.numfmts)
        new_id = max([163] + list(self.custom)) + 1
        el = etree.SubElement(self.numfmts, _q("numFmt"))
        el.set("numFmtId", str(new_id))
        el.set("formatCode", "dd-mmm-yy")
        self.custom[new_id] = "dd-mmm-yy"
        self.numfmts.set("count", str(len(self.numfmts.findall(_q("numFmt")))))
        return new_id

    def date_style(self, s_old: Optional[str]) -> int:
        idx = int(s_old) if s_old not in (None, "") else 0
        if idx in self.cache:
            return self.cache[idx]
        if self.xfs is None:
            return idx
        xfs = self.xfs.findall(_q("xf"))
        if idx >= len(xfs):
            return idx
        xf = xfs[idx]
        if self._is_date(int(xf.get("numFmtId", "0") or 0)):
            res = idx
        else:
            if self.fmt_id is None:
                self.fmt_id = self._add_fmt()
            new = copy.deepcopy(xf)
            new.set("numFmtId", str(self.fmt_id))
            new.set("applyNumberFormat", "1")
            self.xfs.append(new)
            res = len(xfs)
            self.xfs.set("count", str(res + 1))
        self.cache[idx] = res
        return res


def shift_formula(f: str, host_is_target: bool, names: set, at: int, n: int, ext: bool = True) -> str:
    """Shift row numbers >= at by n for references pointing at the target sheet.
    ext=True also extends the END of a range that stops at row at-1."""
    def sh(num: str, lim: int) -> str:
        v = int(num)
        return str(v + n) if v >= lim else num

    def repl(m: re.Match) -> str:
        pre = m.group("pre")
        if pre:
            applies = norm_key(pre[:-1].strip("'").replace("''", "'")) in names
        else:
            applies = host_is_target
        if not applies:
            return m.group(0)
        out = (pre or "") + m.group("c1") + m.group("r1") + sh(m.group("n1"), at)
        if m.group("c2"):
            out += ":" + m.group("c2") + m.group("r2") + sh(m.group("n2"), at - 1 if ext else at)
        return out

    parts = re.split(r'("(?:[^"]|"")*")', f)
    for i in range(0, len(parts), 2):
        parts[i] = _REF_RE.sub(repl, parts[i])
    return "".join(parts)


def _expand_shared(root, ws) -> None:
    for c in root.iter(_q("c")):
        f = c.find(_q("f"))
        if f is not None and f.get("t") == "shared":
            val = ws[c.get("r")].value if ws is not None else None
            if isinstance(val, str) and val.startswith("="):
                for a in ("t", "ref", "si"):
                    f.attrib.pop(a, None)
                f.text = val[1:]


def _shift_formulas(root, host_is_target: bool, names: set, at: int, n: int) -> None:
    for f in root.iter(_q("f")):
        if f.text:
            f.text = shift_formula(f.text, host_is_target, names, at, n)
        if host_is_target and f.get("ref"):
            f.set("ref", shift_formula(f.get("ref"), True, names, at, n))


def _shift_structs(root, at: int, n: int) -> None:
    for tag, attr, ext in (("mergeCell", "ref", False), ("conditionalFormatting", "sqref", True),
                           ("dataValidation", "sqref", True), ("autoFilter", "ref", True),
                           ("dimension", "ref", True), ("hyperlink", "ref", True),
                           ("ignoredError", "sqref", True), ("protectedRange", "sqref", True)):
        for el in root.iter(_q(tag)):
            v = el.get(attr)
            if v:
                el.set(attr, shift_formula(v, True, set(), at, n, ext))
    for tag in ("formula", "formula1", "formula2"):
        for el in root.iter(_q(tag)):
            if el.text:
                el.text = shift_formula(el.text, True, set(), at, n)


def _insert_rows(root, at: int, n: int, src_row: int) -> None:
    """Insert n rows at row `at`, copying style and formulas from src_row (= at-1)."""
    sd = root.find(_q("sheetData"))
    src_el = None
    for row in sd.findall(_q("row")):
        rn = int(row.get("r"))
        if rn == src_row:
            src_el = row
        if rn >= at:
            row.set("r", str(rn + n))
            row.attrib.pop("spans", None)
            for c in row.findall(_q("c")):
                m = re.match(r"([A-Z]+)\d+$", c.get("r"))
                c.set("r", f"{m.group(1)}{rn + n}")
    anchor = next((row for row in sd.findall(_q("row")) if int(row.get("r")) >= at + n), None)
    for k in range(n):
        rn = at + k
        new = etree.Element(_q("row"))
        if src_el is not None:
            for a, v in src_el.attrib.items():
                if a not in ("r", "spans"):
                    new.set(a, v)
            for c in src_el.findall(_q("c")):
                col = re.match(r"([A-Z]+)\d+$", c.get("r")).group(1)
                nc = etree.SubElement(new, _q("c"))
                nc.set("r", f"{col}{rn}")
                if c.get("s") is not None:
                    nc.set("s", c.get("s"))
                f = c.find(_q("f"))
                if f is not None and f.text and f.get("t") in (None, "normal"):
                    try:
                        nf = Translator("=" + f.text, origin=f"{col}{src_row}").translate_formula(f"{col}{rn}")
                        etree.SubElement(nc, _q("f")).text = nf[1:]
                    except Exception:
                        pass
        new.set("r", str(rn))
        if anchor is not None:
            anchor.addprevious(new)
        else:
            sd.append(new)
    mc = root.find(_q("mergeCells"))
    if mc is not None:
        for m in list(mc.findall(_q("mergeCell"))):
            mm = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", m.get("ref") or "")
            if mm and int(mm.group(2)) == src_row == int(mm.group(4)):
                for k in range(n):
                    nm = etree.SubElement(mc, _q("mergeCell"))
                    nm.set("ref", f"{mm.group(1)}{at + k}:{mm.group(3)}{at + k}")
        mc.set("count", str(len(mc.findall(_q("mergeCell")))))


def _shift_drawing(data: bytes, at: int, n: int) -> bytes:
    root = etree.fromstring(data)
    for tag in ("from", "to"):
        for el in root.iter(f"{{{_XDR}}}{tag}"):
            row = el.find(f"{{{_XDR}}}row")
            if row is not None and row.text and int(row.text) >= at - 1:
                row.text = str(int(row.text) + n)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _shift_table(data: bytes, at: int, n: int) -> bytes:
    root = etree.fromstring(data)
    if root.get("ref"):
        root.set("ref", shift_formula(root.get("ref"), True, set(), at, n))
    for el in root.iter(_q("autoFilter")):
        if el.get("ref"):
            el.set("ref", shift_formula(el.get("ref"), True, set(), at, n))
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _sheet_paths(zf: zipfile.ZipFile) -> dict[str, str]:
    wbx = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rid_target = {r.get("Id"): (r.get("Target") or "") for r in rels}
    out: dict[str, str] = {}
    sheets = wbx.find(_q("sheets"))
    for s in (sheets if sheets is not None else []):
        t = rid_target.get(s.get(f"{{{_RNS}}}id"), "")
        t = t.lstrip("/") if t.startswith("/") else "xl/" + t
        out[norm_key(s.get("name"))] = t
    return out


def _tostring(root) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _xml_num(x: float) -> str:
    return str(int(x)) if float(x).is_integer() else repr(float(x))


def _patch_root(root, cells: list, styles: Optional[StyleBook] = None) -> None:
    sd = root.find(_q("sheetData"))
    if sd is None:
        raise ValueError("ไม่พบ sheetData ในชีต")
    rows = {int(r.get("r")): r for r in sd.findall(_q("row")) if r.get("r")}

    def get_row(n: int):
        row = rows.get(n)
        if row is not None:
            row.attrib.pop("spans", None)
            return row
        row = etree.Element(_q("row"))
        row.set("r", str(n))
        later = [k for k in rows if k > n]
        if later:
            rows[min(later)].addprevious(row)
        else:
            sd.append(row)
        rows[n] = row
        return row

    def get_cell(n: int, c: int):
        row = get_row(n)
        ref = f"{get_column_letter(c)}{n}"
        for ce in row.findall(_q("c")):
            if ce.get("r") == ref:
                return ce
        new = etree.Element(_q("c"))
        new.set("r", ref)
        for ce in row.findall(_q("c")):
            m = re.match(r"[A-Z]+", ce.get("r") or "")
            if m and column_index_from_string(m.group(0)) > c:
                ce.addprevious(new)
                return new
        row.append(new)
        return new

    for r, c, v in cells:
        ce = get_cell(r, c)
        if ce.find(_q("f")) is not None:
            continue                                   # never overwrite a formula
        for ch in list(ce):
            ce.remove(ch)
        ce.attrib.pop("t", None)
        if isinstance(v, DateCell):
            if styles is not None:
                ce.set("s", str(styles.date_style(ce.get("s"))))
            etree.SubElement(ce, _q("v")).text = _xml_num(v.serial)
            continue
        ce.set("t", "inlineStr")
        t = etree.SubElement(etree.SubElement(ce, _q("is")), _q("t"))
        t.text = str(v)
        t.set(f"{{{_XML_NS}}}space", "preserve")


def patch_xlsx(master: bytes, cells_by_sheet: dict, ins_by_sheet: dict, wb) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(master))
    try:
        names = set(zin.namelist())
        paths = _sheet_paths(zin)
        key_of_path = {v: k for k, v in paths.items()}
        out_parts: dict[str, bytes] = {}
        roots: dict[str, Any] = {}
        expanded: set[str] = set()
        wbroot = etree.fromstring(zin.read("xl/workbook.xml"))

        def root_of(path: str):
            if path not in roots:
                roots[path] = etree.fromstring(zin.read(path))
            return roots[path]

        def expand(path: str) -> None:
            if path in expanded:
                return
            expanded.add(path)
            ws_o = next((w for w in wb.worksheets if norm_key(w.title) == key_of_path.get(path)), None)
            _expand_shared(root_of(path), ws_o)

        for sheet, ins in ins_by_sheet.items():
            skey = norm_key(sheet)
            spath = paths.get(skey)
            if not spath or spath not in names:
                raise ValueError(f"ไม่พบไฟล์ XML ของชีต {sheet}")
            for after in sorted(ins, reverse=True):
                k = ins[after]
                if k <= 0:
                    continue
                at = after + 1
                tgt = {skey}
                for key, path in paths.items():
                    if path not in names:
                        continue
                    if key != skey and sheet.encode("utf-8") not in zin.read(path):
                        continue
                    expand(path)
                    r = root_of(path)
                    _shift_formulas(r, key == skey, tgt, at, k)
                    if key == skey:
                        _shift_structs(r, at, k)
                        _insert_rows(r, at, k, after)
                for dn in wbroot.iter(_q("definedName")):
                    if dn.text:
                        dn.text = shift_formula(dn.text, False, tgt, at, k)
                rel_path = posixpath.join(posixpath.dirname(spath), "_rels", posixpath.basename(spath) + ".rels")
                if rel_path in names:
                    for rel in etree.fromstring(zin.read(rel_path)):
                        typ = rel.get("Type") or ""
                        tp = rel.get("Target") or ""
                        tp = tp.lstrip("/") if tp.startswith("/") else posixpath.normpath(
                            posixpath.join(posixpath.dirname(spath), tp))
                        if tp not in names:
                            continue
                        if typ.endswith("/drawing"):
                            out_parts[tp] = _shift_drawing(out_parts.get(tp) or zin.read(tp), at, k)
                        elif typ.endswith("/table"):
                            out_parts[tp] = _shift_table(out_parts.get(tp) or zin.read(tp), at, k)

        styles = None
        if "xl/styles.xml" in names and any(isinstance(v, DateCell) for cl in cells_by_sheet.values() for _, _, v in cl):
            styles = StyleBook(zin.read("xl/styles.xml"))
        for sheet, cells in cells_by_sheet.items():
            path = paths.get(norm_key(sheet))
            if not path or path not in names:
                raise ValueError(f"ไม่พบไฟล์ XML ของชีต {sheet}")
            _patch_root(root_of(path), cells, styles)
        for path, r in roots.items():
            out_parts[path] = _tostring(r)
        if styles is not None:
            out_parts["xl/styles.xml"] = _tostring(styles.root)

        calc = wbroot.find(_q("calcPr"))          # make Excel recalculate every formula on open
        if calc is None:
            calc = etree.Element(_q("calcPr"))
            anchor = None
            for tag in ("sheets", "functionGroups", "externalReferences", "definedNames"):
                e = wbroot.find(_q(tag))
                if e is not None:
                    anchor = e
            if anchor is not None:
                anchor.addnext(calc)
            else:
                wbroot.append(calc)
        calc.set("fullCalcOnLoad", "1")
        out_parts["xl/workbook.xml"] = _tostring(wbroot)

        drop: set[str] = set()
        if "xl/calcChain.xml" in names:            # stale after edits; Excel rebuilds it
            drop.add("xl/calcChain.xml")
            ct = etree.fromstring(zin.read("[Content_Types].xml"))
            for el in list(ct):
                if el.get("PartName") == "/xl/calcChain.xml":
                    ct.remove(el)
            out_parts["[Content_Types].xml"] = _tostring(ct)
            rels = etree.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
            for el in list(rels):
                if (el.get("Type") or "").endswith("/calcChain"):
                    rels.remove(el)
            out_parts["xl/_rels/workbook.xml.rels"] = _tostring(rels)

        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename in drop:
                    continue
                data = out_parts.get(item.filename)
                zout.writestr(item, data if data is not None else zin.read(item.filename))
        return out.getvalue()
    finally:
        zin.close()


def build_xlsx(template: bytes, rows: list[dict]) -> tuple[bytes, dict]:
    """rows: dicts with serial / date (ISO or parsable) / form (Service|PM). Rows without serial are skipped."""
    rows = [r for r in rows if clean_serial(r.get("serial"))]
    wb = load_workbook(io.BytesIO(template), keep_links=False)      # read only, never saved
    try:
        plan = analyse_sheet(wb)
        if not plan["ok"]:
            raise ValueError(plan["error"])
        base = date(1904, 1, 1) if getattr(getattr(wb, "epoch", None), "year", 1900) == 1904 else date(1899, 12, 30)
        n = len(rows)
        ins: dict = {}
        inserted = 0
        if plan["free"] is not None and n > plan["free"]:
            inserted = n - plan["free"]
            ins[plan["title"]] = {plan["src"]: inserted}
        cells: list = []
        n_date = n_type = 0
        for i, r in enumerate(rows):
            rn = plan["start"] + i
            cells.append((rn, plan["col"], clean_serial(r.get("serial"))))
            d = parse_date(r.get("date"))
            if d and plan["date_col"]:
                cells.append((rn, plan["date_col"], DateCell((d - base).days)))
                n_date += 1
            if r.get("form") in FORM_OPTIONS and plan["type_col"]:
                cells.append((rn, plan["type_col"], r["form"]))
                n_type += 1
        out = patch_xlsx(template, {plan["title"]: cells} if cells else {}, ins, wb)
    finally:
        wb.close()
    return out, {"first_row": plan["start"], "last_row": plan["start"] + max(n, 1) - 1, "col": plan["col_letter"],
                 "inserted": inserted, "count": n, "dates": n_date, "types": n_type,
                 "no_summary": plan["stop"] is None}


# ----------------------------------------------------------------------------
# Self test
# ----------------------------------------------------------------------------
def _demo_template(date_fmt: str = "dd-mmm-yy") -> bytes:
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    tms = wb.active
    tms.title = "TMS Working Record"
    tms["Z1"] = "=COUNTA('SN tracking form'!$E$7:$E$10)"
    sn = wb.create_sheet(SHEET_SN)
    for c, t in enumerate(["Person in charge (ENG)", "Date", "Hospital Name (ENG)", "Model", "Serial number",
                           "Condition", "Service type"], 1):
        sn.cell(5, c, t)
    sn["A6"] = "(only name)"
    yellow = PatternFill("solid", fgColor="FFFFFFCC")
    for r in range(7, 11):
        sn.cell(r, 5).fill = yellow
        sn.cell(r, 3, f"=IFERROR(INDEX('Data2 (2)'!$L:$L,MATCH($E{r},'Data2 (2)'!$G:$G,0)),\"\")")
        sn.cell(r, 2).number_format = date_fmt
    dv = DataValidation(type="list", formula1='"Service,PM,Installation"', allow_blank=True)
    sn.add_data_validation(dv)
    dv.add("G7:G10")
    sn["D12"], sn["D13"] = "Warranty", "Contract"
    sn["E12"] = "=COUNTIF($F$7:$F$10,D12)"
    d2 = wb.create_sheet(SHEET_DATA2)
    d2.append(["Z", "Type", "Province", "Model name", "Engineer name", "TH - Customer name (Commercial)",
               "Serial number", "ENG - Customer name (Commercial)"])
    for i, s in enumerate(("21CX22507", "21CX22472", "21CX22513", "17CX6407", "2403-055", "14AS4052")):
        d2.append(["", "HD", "Bangkok", "m", "e", "th", s, "Rajthanee Hospital (Ayutthaya)" if i == 0 else "Hospital X"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def run_selftest() -> None:
    assert clean_serial(" 21cx 22507 ") == "21CX22507" and clean_serial("2403-055") == "2403-055"
    assert pick_serial(["21CX22507", "14"]) == "21CX22507" and pick_serial(["", None]) == ""
    assert pick_serial(["21CX2250?", "21CX22507"]) == "21CX22507"
    assert parse_date("6/8/26") == date(2026, 8, 6) and parse_date("2026-08-06") == date(2026, 8, 6)
    assert parse_date("06/08/2569") == date(2026, 8, 6) and parse_date("6/8/69") == date(2026, 8, 6)
    assert parse_date("") is None and parse_date("abc") is None and parse_date("31/2/26") is None
    rows = reports_to_rows(DEMO_PAYLOAD, "a.pdf", 1)
    assert rows[0]["serial"] == "21CX22507" and rows[0]["form"] == "Service"
    assert rows[0]["date"] == "2026-08-06" and rows[0]["customer"].startswith("Rajthanee")
    pm = reports_to_rows({"reports": [{"form_type": "pm", "date": "10/8/26", "candidates": ["21CX22472"]}]}, "b.pdf", 2)
    assert pm[0]["form"] == "PM" and pm[0]["date"] == "2026-08-10" and pm[0]["serial"] == "21CX22472"
    assert reports_to_rows({"reports": []}, "a.pdf", 2)[0]["form"] == ""
    img = Image.new("RGB", (1000, 1400), "white")
    assert Image.open(io.BytesIO(_to_jpeg(img, 0.5))).size == (1000, 700)
    assert Image.open(io.BytesIO(_to_jpeg(img, 1.0))).size == (1000, 1400)
    assert shift_formula("COUNTIF($F$7:$F$10,D12)", True, set(), 11, 2) == "COUNTIF($F$7:$F$12,D14)"
    assert shift_formula("'SN tracking form'!E90", False, {"sn tracking form"}, 11, 2) == "'SN tracking form'!E92"
    assert shift_formula("'Data2 (2)'!$G$2:$G$1995", False, {"sn tracking form"}, 11, 2) == "'Data2 (2)'!$G$2:$G$1995"

    tpl = _demo_template()
    plan = plan_template_impl(tpl)
    assert plan["ok"] and plan["col_letter"] == "E" and plan["start"] == 7 and plan["free"] == 4, plan
    assert plan["src"] == 10 and plan["stop"] == 12, plan
    assert plan["date_letter"] == "B" and plan["type_letter"] == "G", plan
    ref = ref_serials_impl(tpl)
    assert "21CX22507" in ref and ref["21CX22507"]["name"] == "Rajthanee Hospital (Ayutthaya)"
    assert check_serial("21CX22507", ref, list(ref)) == ("พบ", "")
    st_, sug = check_serial("21CX2250", ref, list(ref))
    assert st_ == "ไม่พบ" and sug == "21CX22507", (st_, sug)

    serials = ["21CX22507", "21CX22472", "21CX22513", "17CX6407", "2403-055", "14AS4052"]
    data = [{"serial": s, "date": "2026-08-06" if i != 1 else "", "form": "Service" if i < 4 else "PM"}
            for i, s in enumerate(serials)]
    out, info = build_xlsx(tpl, data)
    assert info["inserted"] == 2 and info["first_row"] == 7 and info["last_row"] == 12, info
    assert info["dates"] == 5 and info["types"] == 6, info
    wb = load_workbook(io.BytesIO(out))
    sn, tms = wb[SHEET_SN], wb["TMS Working Record"]
    assert [sn.cell(r, 5).value for r in range(7, 13)] == serials
    assert sn["B7"].value == datetime(2026, 8, 6) and sn["B8"].value is None and sn["B12"].value == datetime(2026, 8, 6)
    assert is_date_format(sn["B7"].number_format) and is_date_format(sn["B12"].number_format)
    assert [sn.cell(r, 7).value for r in range(7, 13)] == ["Service"] * 4 + ["PM"] * 2
    assert sn["D14"].value == "Warranty" and sn["D15"].value == "Contract"
    assert sn["E14"].value == "=COUNTIF($F$7:$F$12,D14)", sn["E14"].value
    assert tms["Z1"].value == "=COUNTA('SN tracking form'!$E$7:$E$12)", tms["Z1"].value
    assert "$E11" in sn["C11"].value and "$E12" in sn["C12"].value and "$E7" in sn["C7"].value
    assert sn["E11"].fill.fill_type == "solid" and sn["E12"].fill.fill_type == "solid"
    dvs = sn.data_validations.dataValidation
    assert len(dvs) == 1 and str(dvs[0].sqref) == "G7:G12" and "Service" in dvs[0].formula1, [str(d.sqref) for d in dvs]
    # fewer serials than free rows: no insertion
    out2, info2 = build_xlsx(tpl, data[:3])
    assert info2["inserted"] == 0
    assert load_workbook(io.BytesIO(out2))[SHEET_SN]["E14"].value is None
    # template whose Date cells are "General": app gives them a date format
    out3, _ = build_xlsx(_demo_template("General"), data)
    sn3 = load_workbook(io.BytesIO(out3))[SHEET_SN]
    assert sn3["B7"].value == datetime(2026, 8, 6) and is_date_format(sn3["B7"].number_format)
    assert sn3["B12"].value == datetime(2026, 8, 6) and is_date_format(sn3["B12"].number_format)
    print("SELFTEST PASSED")


# ----------------------------------------------------------------------------
# Streamlit UI
# ----------------------------------------------------------------------------
def _stretch(fn, *args, **kwargs):
    try:
        return fn(*args, width="stretch", **kwargs)
    except Exception:
        return fn(*args, use_container_width=True, **kwargs)


def _secret(*names: str) -> str:
    for name in names:
        try:
            v = str(st.secrets.get(name, ""))
        except Exception:
            v = ""
        v = v or os.environ.get(name, "")
        if v:
            return v
    return ""


@st.cache_data(show_spinner=False)
def ref_serials(data: bytes) -> dict:
    return ref_serials_impl(data)


@st.cache_data(show_spinner=False)
def plan_template(data: bytes) -> dict:
    return plan_template_impl(data)


def sidebar() -> dict:
    sb = st.sidebar
    sb.header("ตั้งค่า AI")
    provider = sb.selectbox("Provider", [PROVIDER_GEMINI, PROVIDER_CLAUDE, PROVIDER_OPENAI, PROVIDER_DEMO])
    names = (("GEMINI_API_KEY", "GOOGLE_API_KEY") if provider == PROVIDER_GEMINI
             else ("ANTHROPIC_API_KEY",) if provider == PROVIDER_CLAUDE else ("OPENAI_API_KEY",))
    key = sb.text_input("API key", value=_secret(*names), type="password", disabled=provider == PROVIDER_DEMO)
    model = sb.text_input("Model", value=DEFAULT_MODELS[provider], key=f"model_{provider}")
    sb.caption("ใส่ key ไว้ใน Settings > Secrets แล้วปล่อยช่องนี้ว่างได้")
    sb.header("ความเร็ว")
    crop_label = sb.selectbox("ส่วนของหน้าที่ให้ AI อ่าน", list(CROP_OPTIONS))
    sb.caption("ถ้าอ่านส่วนบนแล้วไม่เจอ Serial แอปจะลองอ่านทั้งหน้าให้อัตโนมัติ")
    pause = sb.number_input("หน่วงเวลาระหว่างหน้า (วินาที)", min_value=0.0, max_value=30.0,
                            value=GEMINI_PAUSE if provider == PROVIDER_GEMINI else 0.0, step=0.5,
                            key=f"pause_{provider}")
    sb.caption("Gemini ฟรีควรใช้ 4 วินาที ถ้าผูกบัตรแล้วตั้งเป็น 0 ได้")
    return {"provider": provider, "key": key, "model": model, "crop": CROP_OPTIONS[crop_label], "pause": pause}


def init_state() -> None:
    if "rows" not in st.session_state:
        store = load_store()
        st.session_state.rows = store["rows"]
        st.session_state.live = store["rows"]
        st.session_state.done = set(store["done"])
        st.session_state.ver = 0
        st.session_state.upver = 0


def main() -> None:
    st.set_page_config(page_title="Serial Number Filler", page_icon="🔢", layout="wide")
    init_state()
    cfg = sidebar()
    ss = st.session_state

    st.title("กรอก Serial number / Date / Service type ลงไฟล์ Excel")
    st.caption("1) อัปโหลดไฟล์ Blank  →  2) อัปโหลดใบงาน  →  3) ตรวจ/แก้ไข  →  4) Export")
    if fitz is None:
        st.error("ไม่พบ PyMuPDF (ยังอ่านไฟล์รูปภาพได้ แต่อ่าน PDF ไม่ได้)")
    if ss.get("last_msg"):
        st.success(ss.pop("last_msg"))
    for e in ss.pop("last_errors", []):
        st.error(e)

    # ---- Step 1 -----------------------------------------------------------
    st.header("1) อัปโหลดไฟล์ Blank file (.xlsx)")
    up = st.file_uploader("ไฟล์ที่ผูกสูตรไว้แล้ว", type=["xlsx"], key="tpl_up")
    if up is not None:
        data = up.getvalue()
        cur = TEMPLATE_PATH.read_bytes() if TEMPLATE_PATH.exists() else b""
        if hashlib.sha1(data).digest() != hashlib.sha1(cur).digest():
            DATA_DIR.mkdir(parents=True, exist_ok=True)
            TEMPLATE_PATH.write_bytes(data)
            ss.pop("out", None)
    template = TEMPLATE_PATH.read_bytes() if TEMPLATE_PATH.exists() else None
    plan = None
    if template is None:
        st.info("ยังไม่มีไฟล์ Blank")
    else:
        plan = plan_template(template)
        if not plan["ok"]:
            st.error(plan["error"])
        else:
            msg = (f"พบชีต **{plan['title']}** | Serial number = คอลัมน์ **{plan['col_letter']}** | "
                   f"Date = **{plan['date_letter'] or 'ไม่พบ'}** | Service type = **{plan['type_letter'] or 'ไม่พบ'}** | "
                   f"เริ่มลงที่แถว **{plan['start']}**")
            if plan["free"] is not None:
                msg += f" | ช่องว่างที่เหลือ {plan['free']} ช่อง (ถ้าไม่พอจะแทรกแถวให้อัตโนมัติ)"
            st.success(msg)
            if not plan["date_col"]:
                st.warning("ไม่พบหัวคอลัมน์ 'Date' จะไม่ลงวันที่")
            if not plan["type_col"]:
                st.warning("ไม่พบหัวคอลัมน์ 'Service type' จะไม่ลงชนิดงาน")
            if plan["existing"]:
                st.warning(f"ในไฟล์มี Serial อยู่แล้ว {plan['existing']} แถว ข้อมูลใหม่จะต่อท้ายจากแถวสุดท้าย")
            if plan["stop"] is None:
                st.warning("ไม่พบตารางสรุป (Warranty / Contract) ท้ายชีต จะไม่แทรกแถวให้ ถ้าเลขเกินช่องที่จัดรูปแบบไว้ "
                           "แถวส่วนเกินจะไม่มีสี/เส้นตาราง")

    # ---- Step 2 -----------------------------------------------------------
    st.header("2) อัปโหลดใบงาน (Service report / PM report)")
    files = st.file_uploader("PDF หลายหน้า หรือรูปภาพ เลือกได้หลายไฟล์ และอัปโหลดเพิ่มทีหลังได้",
                             type=["pdf", "png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp"],
                             accept_multiple_files=True, key=f"up_{ss.upver}")
    if st.button("อ่านใบงาน", type="primary", disabled=not files):
        rows = [dict(r) for r in ss.live]
        done = set(ss.done)
        client = VisionClient(cfg["provider"], cfg["key"], cfg["model"], cfg["pause"])
        bar, info = st.progress(0.0), st.empty()

        def cb(p: float, text: str) -> None:
            bar.progress(min(max(p, 0.0), 1.0))
            info.write(f"กำลังอ่าน: {text}")

        try:
            res = process_files([(f.name, f.getvalue()) for f in files], client, rows, done, cfg["crop"], cb)
        finally:
            ss.rows, ss.live, ss.done = rows, rows, done
            save_store(rows, done)
        errs = list(res["errors"])
        if res["stopped"]:
            errs.append("หยุดอ่านระหว่างทาง: " + res["stopped"] +
                        " (หน้าที่อ่านเสร็จแล้วถูกเก็บไว้ กดอ่านใหม่ภายหลังเพื่ออ่านต่อ)")
        ss.last_msg = f"อ่านแล้ว {res['pages']} หน้า ได้ {res['added']} แถว ข้ามหน้าที่เคยอ่าน {res['skipped']} หน้า"
        ss.last_errors = errs
        ss.ver += 1
        ss.upver += 1
        ss.pop("out", None)
        st.rerun()

    # ---- Step 3 -----------------------------------------------------------
    st.header("3) ตรวจและแก้ไขข้อมูล")
    st.caption("แก้ในตารางได้โดยตรง เพิ่มแถวที่ท้ายตาราง หรือเลือกแถวแล้วกด Delete เพื่อลบ "
               "วันที่พิมพ์เป็น YYYY-MM-DD หรือ วัน/เดือน/ปี ก็ได้ ลำดับแถวในตารางคือลำดับที่จะลงใน Excel")
    df = pd.DataFrame(ss.rows, columns=COLS)
    df = df.astype(object).where(df.notna(), None)
    df["form"] = df["form"].apply(lambda v: v if v in FORM_OPTIONS else None)
    edited = _stretch(
        st.data_editor, df, key=f"ed_{ss.ver}", num_rows="dynamic", hide_index=True,
        disabled=["source", "page"],
        column_config={
            "serial": st.column_config.TextColumn("Serial number"),
            "date": st.column_config.TextColumn("Date (YYYY-MM-DD)"),
            "customer": st.column_config.TextColumn("ชื่อลูกค้าที่อ่านได้ (ไม่ลง Excel)"),
            "form": st.column_config.SelectboxColumn("Service type", options=FORM_OPTIONS),
            "source": st.column_config.TextColumn("ไฟล์"),
            "page": st.column_config.TextColumn("หน้า"),
        })
    live = df_to_rows(edited)
    ss.live = live
    save_store(live, ss.done)

    ref = ref_serials(template) if template else {}
    ref_keys = list(ref)
    results = [check_serial(r["serial"], ref, ref_keys) for r in live]
    n_ok = sum(1 for s, _ in results if s == "พบ")
    bad = [i for i, (s, _) in enumerate(results) if s in ("ไม่พบ", "ว่าง")]
    dup = [s for s, c in Counter(r["serial"] for r in live if r["serial"]).items() if c > 1]
    q = sum(1 for r in live if "?" in r["serial"])
    with_serial = [r for r in live if r["serial"]]
    n_bad_date = sum(1 for r in with_serial if parse_date(r["date"]) is None)
    n_no_form = sum(1 for r in with_serial if not r["form"])
    st.write(f"ทั้งหมด **{len(live)}** แถว" + (f" | พบใน Data2 {n_ok} | ต้องตรวจ {len(bad)}" if ref else ""))
    if q:
        st.warning(f"มี {q} แถวที่มีเครื่องหมาย ? (ตัวอักษรที่ AI อ่านไม่ออก) กรุณาแก้ในตาราง")
    if n_bad_date:
        st.warning(f"มี {n_bad_date} แถวที่วันที่ว่างหรืออ่านไม่ได้ แถวเหล่านี้จะไม่ลงวันที่ใน Excel (กรอกในตารางได้)")
    if n_no_form:
        st.info(f"มี {n_no_form} แถวที่ไม่ได้ระบุชนิดใบงาน ช่อง Service type ใน Excel จะเว้นว่างให้เลือกเอง")
    if dup:
        st.info("Serial ซ้ำ: " + ", ".join(dup[:15]) + " (ปกติถ้าเครื่องเดียวมีทั้งใบ Service และ PM)")
    if ref and live:
        with st.expander("เทียบกับ Data2 (ชื่อ รพ. ที่อ่านได้ กับชื่อใน Data2 ตาม Serial)", expanded=bool(bad)):
            table = []
            for i, r in enumerate(live):
                s, sug = results[i]
                d = parse_date(r["date"])
                table.append({"แถวที่": i + 1, "Serial": r["serial"], "สถานะ Serial": s,
                              "ค่าใกล้เคียงใน Data2": sug,
                              "รพ. ใน Data2": ref[r["serial"]]["name"] if r["serial"] in ref else "",
                              "ชื่อที่อ่านได้": r["customer"],
                              "วันที่": d.isoformat() if d else ("ผิดรูปแบบ: " + r["date"] if r["date"] else "ว่าง"),
                              "ชนิด": r["form"], "ไฟล์": r["source"], "หน้า": r["page"]})
            _stretch(st.dataframe, pd.DataFrame(table), hide_index=True)
            if any(results[i][1] for i in bad) and st.button("ใช้ค่าที่แนะนำกับทุกแถวที่มีค่าใกล้เคียง"):
                new = [dict(r) for r in live]
                for i in bad:
                    if results[i][1]:
                        new[i]["serial"] = clean_serial(results[i][1])
                ss.rows = ss.live = new
                save_store(new, ss.done)
                ss.ver += 1
                ss.pop("out", None)
                st.rerun()

    with st.expander("ล้างข้อมูลทั้งหมด"):
        if st.checkbox("ยืนยันล้างข้อมูลที่อ่านไว้และประวัติหน้าที่อ่านแล้ว") and st.button("ล้างข้อมูล"):
            ss.rows, ss.live, ss.done = [], [], set()
            save_store([], set())
            ss.ver += 1
            ss.upver += 1
            ss.pop("out", None)
            st.rerun()

    # ---- Step 4 -----------------------------------------------------------
    st.header("4) Export")
    if plan and plan.get("ok"):
        need = max(len(with_serial) - plan["free"], 0) if plan["free"] is not None else 0
        st.caption(f"จะลง {len(with_serial)} แถว: Serial ที่ช่อง {plan['col_letter']}{plan['start']} เป็นต้นไป"
                   f" | Date ที่คอลัมน์ {plan['date_letter'] or '-'} | Service type ที่คอลัมน์ {plan['type_letter'] or '-'}"
                   + (f" | แทรกแถวเพิ่ม {need} แถว" if need else ""))
    if st.button("สร้างไฟล์ Excel", type="primary", disabled=template is None or not with_serial):
        try:
            data, info = build_xlsx(template, live)
            ss.out = (data, info, f"SN_filled_{datetime.now():%Y%m%d_%H%M}.xlsx")
        except Exception as e:
            ss.pop("out", None)
            st.error(f"สร้างไฟล์ไม่สำเร็จ: {e}")
    if ss.get("out"):
        data, info, fname = ss.out
        st.success(f"สร้างไฟล์แล้ว: ลง {info['count']} แถว ที่แถว {info['first_row']}–{info['last_row']} "
                   f"(วันที่ {info['dates']} ช่อง, Service type {info['types']} ช่อง)"
                   + (f" | แทรกแถวเพิ่ม {info['inserted']} แถว" if info["inserted"] else ""))
        st.download_button("ดาวน์โหลด Excel", data=data, file_name=fname,
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        st.caption("ถ้าแก้ตารางหลังสร้างไฟล์ ให้กดสร้างไฟล์ใหม่อีกครั้ง")


if __name__ == "__main__" and "--selftest" in sys.argv:
    run_selftest()
else:
    main()
